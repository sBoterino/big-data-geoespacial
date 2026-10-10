from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "informe_assets"
OUTPUT_DOCX = DOCS / "informe_tecnico_big_data_geoespacial.docx"
ARCH = ASSETS / "arquitectura.png"
BENCH = DOCS / "evidencias" / "bench_tiempo.png"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
ORANGE = "ED7D31"
LIGHT_GRAY = "D9E1E8"
MID_GRAY = "65727E"
BLACK = "000000"
WHITE = "FFFFFF"


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, subtitle, fill, outline=NAVY):
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline="#" + outline, width=4)
    x1, y1, x2, y2 = xy
    title_font = font(31, True)
    sub_font = font(22)
    title_box = draw.textbbox((0, 0), title, font=title_font)
    title_w = title_box[2] - title_box[0]
    draw.text(((x1 + x2 - title_w) / 2, y1 + 22), title, font=title_font, fill="#" + BLACK)
    lines = subtitle.split("\n")
    y = y1 + 67
    for line in lines:
        box = draw.textbbox((0, 0), line, font=sub_font)
        w = box[2] - box[0]
        draw.text(((x1 + x2 - w) / 2, y), line, font=sub_font, fill="#33404A")
        y += 27


def arrow(draw, start, end, label=None):
    draw.line([start, end], fill="#" + MID_GRAY, width=6)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 12), (x2 - 18 * direction, y2 + 12)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 12, y2 - 18 * direction), (x2 + 12, y2 - 18 * direction)]
    draw.polygon(points, fill="#" + MID_GRAY)
    if label:
        label_font = font(18)
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        box = draw.textbbox((0, 0), label, font=label_font)
        draw.rectangle((mx - 6, my - 18, mx + (box[2] - box[0]) + 6, my + 8), fill="white")
        draw.text((mx, my - 16), label, font=label_font, fill="#" + MID_GRAY)


def build_architecture():
    ASSETS.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 35), "Arquitectura y flujos del sistema", font=font(44, True), fill="#" + BLACK)
    draw.text((70, 91), "Flujo de datos", font=font(25, True), fill="#" + NAVY)

    boxes = [
        ((70, 145, 330, 275), "Kaggle", "CSV original\n1.972.121 filas", "#F2F2F2"),
        ((390, 145, 650, 275), "Dask", "limpieza R1-R6\n2 workers", "#DCE6F1"),
        ((710, 145, 970, 275), "MongoDB", "GeoJSON Point\nindice 2dsphere", "#E2F0D9"),
        ((1030, 145, 1290, 275), "Spark", "grilla y series\ntemporales", "#FCE4D6"),
        ((1350, 145, 1610, 275), "Flask", "API REST\nconsultas y resultados", "#E4DFEC"),
    ]
    for xy, title, subtitle, fill in boxes:
        rounded_box(draw, xy, title, subtitle, fill)
    for left, right in zip(boxes[:-1], boxes[1:]):
        arrow(draw, (left[0][2], 210), (right[0][0], 210))
    rounded_box(draw, (1510, 315, 1750, 425), "Cliente", "PowerShell\no navegador", "#FFF2CC")
    arrow(draw, (1610, 275), (1630, 315), "JSON")

    draw.text((70, 385), "Flujo de integracion y despliegue", font=font(25, True), fill="#" + NAVY)
    ci_boxes = [
        ((70, 450, 330, 580), "GitHub", "rama, PR\ny merge", "#F2F2F2"),
        ((390, 450, 650, 580), "smee", "webhook saliente\nsin exponer Jenkins", "#FFF2CC"),
        ((710, 450, 970, 580), "Jenkins", "build, pruebas\ny staging", "#DCE6F1"),
        ((1030, 450, 1290, 580), "Docker", "Compose y red\nbdgeo-net", "#E2F0D9"),
        ((1350, 450, 1610, 580), "Produccion", "API versionada\nsolo si todo pasa", "#E4DFEC"),
    ]
    for xy, title, subtitle, fill in ci_boxes:
        rounded_box(draw, xy, title, subtitle, fill)
    for left, right in zip(ci_boxes[:-1], ci_boxes[1:]):
        arrow(draw, (left[0][2], 515), (right[0][0], 515))

    draw.rounded_rectangle((350, 115, 1645, 635), radius=30, outline="#A6A6A6", width=3)
    draw.text((1175, 610), "Servicios internos en Docker Compose", font=font(19, True), fill="#6B6B6B")
    draw.text((70, 680), "Persistencia local: volumen MongoDB, volumen de datos Kaggle y volumen Jenkins.", font=font(22), fill="#" + MID_GRAY)
    img.save(ARCH, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=10.5, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pagina ")
    set_run_font(run, size=8, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_text(doc, text, bold_lead=None, size=10.5, after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=size, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_run_font(body, size=size)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size)
    return p


def add_bullets(doc, items, size=10.0, after=2):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.03
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(0 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 11.5, bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            if widths:
                cell.width = Inches(widths[cidx])
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


def add_picture_with_alt(paragraph, path, width, alt_text):
    shape = paragraph.add_run().add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return shape


def build_document():
    build_architecture()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor.from_string(BLACK)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("Procesamiento y consulta de datos geoespaciales")
    set_run_font(hr, size=8, color=MID_GRAY)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    # Page 1 - Cover and executive summary
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Procesamiento y consulta de datos geoespaciales")
    set_run_font(r, size=25, bold=True)
    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("Informe tecnico del sistema distribuido y su despliegue continuo")
    set_run_font(r, size=13, color=MID_GRAY)

    add_text(doc, "Juan Guillermo Echeverri", size=11.5, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_text(doc, "Sebastian Botero Velasquez", size=11.5, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_text(doc, "Santiago Villamizar Mejia", size=11.5, after=18, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_text(doc, "Octubre de 2026", size=10.5, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_text(doc, "Repositorio: https://github.com/sBoterino/big-data-geoespacial", size=9.5, after=30, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "Resumen ejecutivo", 1)
    add_text(doc, "El sistema procesa el dataset NYC Motor Vehicle Collisions - Crashes con 1.972.121 registros. Dask ejecuta la limpieza distribuida, MongoDB almacena 1.741.828 documentos geoespaciales, Spark calcula agregaciones espaciales y temporales, Flask publica las consultas y Jenkins automatiza pruebas, staging y despliegue. La carga completa se reprodujo en un segundo computador y el pipeline impide desplegar cuando una prueba falla.")
    add_text(doc, "Resultado principal. La arquitectura satisface el procesamiento de mas de un millon de registros, conserva GeoJSON valido y resultados verificables, y puede levantarse desde GitHub con Docker Compose y secretos locales.", bold_lead="Resultado principal.")

    # Page 2 - Architecture
    page_break(doc)
    add_heading(doc, "1 Arquitectura del sistema", 1)
    add_text(doc, "La solucion separa el flujo de datos del flujo de entrega. Los servicios se comunican por la red interna bdgeo-net mediante nombres de servicio; los puertos publicados se usan solo desde el computador anfitrion. MongoDB, los datos descargados y Jenkins persisten en volumenes Docker.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_picture_with_alt(
        p,
        ARCH,
        Inches(6.95),
        "Diagrama de arquitectura con flujo Kaggle, Dask, MongoDB, Spark, Flask y CI/CD con GitHub, smee, Jenkins y Docker.",
    )
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    rr = cap.add_run("Figura 1. Arquitectura, flujo de datos y flujo CI/CD.")
    set_run_font(rr, size=8.5, color=MID_GRAY)
    add_table(
        doc,
        ["Componente", "Responsabilidad", "Despliegue"],
        [
            ["Dask", "Descarga, limpieza R1-R6 e ingesta por particiones", "scheduler y 2 workers"],
            ["MongoDB", "Documentos, validacion GeoJSON e indices", "servicio con volumen"],
            ["Spark", "Agregaciones por grilla y tiempo", "master y worker"],
            ["Flask", "API REST, validacion y tiempos", "Gunicorn versionado"],
            ["Jenkins", "Build, pruebas, staging y despliegue", "perfil ci"],
        ],
        widths=[1.15, 4.25, 1.45],
        font_size=8.2,
    )

    # Page 3 - Data and cleaning
    page_break(doc)
    add_heading(doc, "2 Datos, limpieza e ingesta", 1)
    add_text(doc, "Se selecciono NYC Motor Vehicle Collisions - Crashes porque aporta latitud, longitud, fecha, hora y defectos reales de calidad. La copia verificada el 1 de octubre de 2026 contiene 1.972.121 filas, 29 columnas y 420.704.526 bytes. Dask proceso seis particiones con dos workers y registro cada descarte en la primera regla incumplida.")
    add_table(
        doc,
        ["Regla", "Criterio", "Descartados"],
        [
            ["R1", "Coordenadas nulas o no numericas", "226.028"],
            ["R2", "Punto centinela (0,0)", "4.115"],
            ["R3", "Fuera de rangos WGS84", "106"],
            ["R4", "Valida globalmente pero fuera de NYC", "44"],
            ["R5", "Fecha u hora invalida", "0"],
            ["R6", "COLLISION_ID duplicado", "0"],
            ["Final", "Documentos validos insertados", "1.741.828"],
        ],
        widths=[0.65, 5.15, 1.05],
        font_size=8.8,
    )
    add_text(doc, "Control de integridad. Los 230.293 descartes y los 1.741.828 documentos finales suman exactamente las 1.972.121 filas iniciales. La carga completa tardo 57,435 s en el equipo de desarrollo.", bold_lead="Control de integridad.")
    add_heading(doc, "Idempotencia y persistencia", 2)
    add_text(doc, "La coleccion ingesta_meta registra dataset, limite de filas, estado y total final. Una segunda ejecucion compara esos metadatos con el conteo real y muestra Ingesta omitida si coinciden. FORCE_RELOAD permite una recarga explicita. El CSV y el Parquet permanecen en un volumen compartido, por lo que los builds no descargan de nuevo el archivo.")
    add_heading(doc, "Modelo de documento", 2)
    add_text(doc, "COLLISION_ID se usa como _id; fecha se guarda como tipo BSON Date; location es un GeoJSON Point en orden [longitud, latitud]. MongoDB aplica un esquema de validacion y mantiene los indices _id_, fecha_1 y location_2dsphere.")

    # Page 4 - Mongo and API
    page_break(doc)
    add_heading(doc, "3 Consultas geoespaciales y API", 1)
    add_text(doc, "La API Flask valida tipos, rangos y geometria antes de consultar MongoDB. Todas las operaciones aceptan limites, responden JSON, devuelven tiempo_ms y rechazan entradas invalidas con HTTP 400. El indice 2dsphere habilita las busquedas sobre la superficie terrestre.")
    add_table(
        doc,
        ["Metodo y ruta", "Operacion", "Uso"],
        [
            ["GET /health", "Ping a MongoDB y version", "Estado y evidencia del despliegue"],
            ["GET /near", "$near", "Vecinos ordenados por distancia"],
            ["POST /within", "$geoWithin", "Eventos dentro de un Polygon GeoJSON"],
            ["GET /geonear", "$geoNear", "Distancia y agrupacion opcional"],
            ["GET /spark-results", "Lectura controlada", "Catalogo de resultados Spark"],
            ["GET /spark-results/<nombre>", "Consulta de coleccion", "Hotspots y series temporales"],
        ],
        widths=[1.85, 1.45, 3.55],
        font_size=8.5,
    )
    add_heading(doc, "Pruebas sobre 1.741.828 documentos", 2)
    add_table(
        doc,
        ["Consulta", "Resultado", "Promedio caliente"],
        [
            ["$near", "100 documentos", "13,493 ms"],
            ["$geoWithin", "100 documentos", "99,531 ms"],
            ["$geoNear simple", "100 documentos", "182,280 ms"],
            ["$geoNear agrupado", "5 grupos", "1.786,176 ms"],
        ],
        widths=[2.2, 2.25, 2.4],
        font_size=9.0,
    )
    add_text(doc, "Las mediciones calientes excluyen la primera ejecucion. $geoNear agrupado es mas costoso porque calcula distancias para todos los eventos del radio antes de agrupar; la variante simple puede detenerse al completar el limite. /within recibe cualquier poligono valido, por lo que cambiar la zona no exige desplegar codigo.")

    # Page 5 - Spark
    page_break(doc)
    add_heading(doc, "4 Procesamiento y verificacion con Spark", 1)
    add_text(doc, "Spark lee desde MongoDB mediante el MongoDB Spark Connector. Un pipeline de proyeccion elimina campos innecesarios antes de transferir datos. La grilla regular usa celdas de 0,005 grados, aproximadamente 555 por 420 m en Nueva York, y puede modificarse con el parametro --celda.")
    add_table(
        doc,
        ["Salida", "Contenido", "Control"],
        [
            ["spark_grilla", "n, heridos, muertos y poligono por celda", "suma n = total"],
            ["spark_hotspots", "20 celdas de mayor concentracion", "ranking reproducible"],
            ["spark_por_hora", "conteo por hora", "suma n = total"],
            ["spark_por_dia_semana", "conteo lunes a domingo", "suma n = total"],
            ["spark_por_mes", "conteo mensual", "suma n = total"],
            ["spark_hora_borough", "hora por borough", "consistencia interna"],
        ],
        widths=[2.1, 3.45, 1.3],
        font_size=8.4,
    )
    add_heading(doc, "Verificacion independiente", 2)
    add_text(doc, "verificar_resultados.py exige que grilla, hora, dia y mes sumen el total de documentos y vuelve a contar el hotspot principal con $geoWithin en MongoDB, con tolerancia de 0,5 %. En la semilla determinista, Spark obtuvo 57 y MongoDB 57. Sobre datos reales, Spark proceso 1.741.828 registros en 3.246 celdas; el hotspot principal tuvo 5.336 en Spark y 5.337 en MongoDB, dentro del margen permitido.")
    add_heading(doc, "Semilla determinista", 2)
    add_text(doc, "La coleccion eventos_semilla contiene 300 puntos de tres zonas de Nueva York. Los controles 120, 80 y 100 verifican conectividad, orden [longitud, latitud] y consultas espaciales sin depender de la carga completa. Cada build escribe resultados spark_semilla_* para no alterar las colecciones reales que consume la API.")

    # Page 6 - CI/CD
    page_break(doc)
    add_heading(doc, "5 Integracion, pruebas y despliegue continuo", 1)
    add_text(doc, "El trabajo se integra mediante ramas y pull requests. Un merge a main genera un push; GitHub lo envia al canal smee y el cliente local lo reenvia a Jenkins por la red interna. Jenkins usa el Docker del host, pero no se publica en Internet.")
    add_table(
        doc,
        ["Etapa", "Comprobacion principal"],
        [
            ["Checkout", "Commit, autor y Jenkinsfile versionado"],
            ["Construir imagenes", "API, ingesta, Spark y servicios Compose"],
            ["Pruebas unitarias", "30 API, 5 ingesta y 4 Spark"],
            ["Kaggle", "Credencial y disponibilidad del dataset"],
            ["Servicios e ingesta", "Healthchecks e idempotencia en main"],
            ["Integracion", "2 workers Dask y Spark-MongoDB"],
            ["Procesamiento Spark", "Semilla siempre; datos reales bajo condicion"],
            ["Staging", "Imagen candidata y smoke tests 120/120/120"],
            ["Despliegue", "Solo main y solo despues de aprobar todo"],
        ],
        widths=[2.05, 4.8],
        font_size=8.2,
    )
    add_text(doc, "Bloqueo demostrado. Al ejecutar FORZAR_FALLO=true, pytest produce un fallo intencional, Jenkins omite las etapas posteriores y /health conserva la version anterior. En los jobs de rama, staging usa eventos_semilla y el despliegue se omite, lo que evita modificar produccion durante una validacion.", bold_lead="Bloqueo demostrado.")
    add_text(doc, "Credenciales. MongoDB se inyecta desde mongo-root y Kaggle desde kaggle-api-token. En desarrollo se usa un .env ignorado por Git. Ningun token, contraseña, CSV o volumen se versiona.", bold_lead="Credenciales.")

    # Page 7 - Benchmark
    page_break(doc)
    add_heading(doc, "6 Analisis comparativo Dask y Spark", 1)
    add_text(doc, "Ambos motores ejecutaron la misma operacion: leer el mismo Parquet, asignar cada punto a una celda de 0,005 grados, contar por celda y obtener el top 20. Se probaron uno y dos workers, con 1.741.828 filas y con una copia fisica de 17.418.280 filas. Cada combinacion tuvo calentamiento y tres repeticiones validas.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    add_picture_with_alt(
        p,
        BENCH,
        Inches(6.7),
        "Grafica de tiempos de calculo de Dask y Spark para 1,7 millones y 17,4 millones de filas con uno y dos workers.",
    )
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    rr = cap.add_run("Figura 2. Tiempo de calculo; media de tres repeticiones sin calentamiento.")
    set_run_font(rr, size=8.2, color=MID_GRAY)
    add_table(
        doc,
        ["Datos", "Dask A / B", "Spark A / B", "Memoria observada Dask / Spark"],
        [
            ["x1", "0,30 / 0,20 s", "3,80 / 4,65 s", "425-612 / 1.077-1.657 MiB"],
            ["x10", "2,36 / 1,39 s", "4,51 / 5,12 s", "491-698 / 1.121-1.816 MiB"],
        ],
        widths=[0.65, 1.55, 1.55, 3.1],
        font_size=8.3,
    )
    add_text(doc, "Dask fue mas rapido en las cuatro combinaciones. En la configuracion A, la ventaja bajo de 12,7 veces en x1 a 1,9 veces en x10: el costo fijo de Spark pesa menos al crecer el volumen. Dask mejoro con dos workers; Spark empeoro porque ambos workers compartian el mismo PC y añadieron comunicacion entre JVM. Los dos motores produjeron la misma celda top: n=5.336 en x1 y 53.360 en x10.")
    add_text(doc, "Limitaciones. La prueba se ejecuto en un solo computador, la memoria se muestreo cada aproximadamente 2 s y x10 repite filas. No se concluye que Dask sea universalmente superior ni que Spark vaya a superarlo; solo se reportan las condiciones medidas.", bold_lead="Limitaciones.", size=9.4)

    # Page 8 - Decisions and reproducibility
    page_break(doc)
    add_heading(doc, "7 Decisiones y reproducibilidad", 1)
    add_table(
        doc,
        ["Decision", "Motivo y efecto"],
        [
            ["MongoDB local", "Cumple Compose, evita limite de Atlas y reduce dependencia de red"],
            ["CI/CD desde el inicio", "Reduce el riesgo de integrar Jenkins al final"],
            ["Ingesta idempotente", "Evita descargar y recargar casi dos millones de filas"],
            ["Conector en imagen Spark", "Version fija y builds repetibles sin descarga por ejecucion"],
            ["Semilla de 300 puntos", "Pruebas rapidas y resultados espaciales conocidos"],
            ["Staging antes de promover", "La misma imagen probada llega a produccion"],
            ["smee en vez de ngrok", "Webhook saliente sin exponer Jenkins"],
            ["Parquet fisico x10", "Comparacion de volumen sin reutilizacion asimetrica de tareas"],
        ],
        widths=[2.0, 4.85],
        font_size=8.3,
    )
    add_heading(doc, "Reproduccion en un segundo computador", 2)
    add_text(doc, "Juan Guillermo clono main en Windows con 16 GB de RAM, creo sus propios secretos y construyo los servicios sin copiar carpetas, imagenes, volumenes ni datos. Verifico dos workers Dask, el conector Spark-MongoDB, la carga completa, la segunda ingesta omitida, RESULTADO: OK de Spark, /near y hotspots. La incidencia inicial fue Docker Desktop instalado pero con el motor apagado; abrirlo resolvio el problema.")
    add_heading(doc, "Arranque desde cero", 2)
    add_bullets(doc, [
        "Clonar el repositorio y copiar .env.example a .env.",
        "Definir contraseña MongoDB, token Kaggle y limites de memoria.",
        "Ejecutar docker compose up -d --build y comprobar /health.",
        "Configurar Jenkins con mongo-root, kaggle-api-token y el job bdgeo-main.",
        "Activar smee, hacer merge mediante PR y observar el pipeline.",
    ], size=9.4)

    # Page 9 - Conclusions and references
    page_break(doc)
    add_heading(doc, "8 Conclusiones", 1)
    add_bullets(doc, [
        "El sistema proceso 1.972.121 filas, justifico 230.293 descartes y almaceno 1.741.828 documentos GeoJSON con indice 2dsphere.",
        "Dask resolvio la ingesta con dos workers e idempotencia; Spark genero agregaciones espaciales y temporales verificadas contra MongoDB.",
        "La API ofrece consultas parametrizadas, validacion y tiempos sobre datos reales, ademas de resultados precalculados de Spark.",
        "El pipeline automatiza pruebas, integracion, staging y despliegue; una prueba fallida conserva la version anterior.",
        "En el hardware medido, Dask fue mas rapido y uso menos memoria observada; la diferencia se redujo con diez veces mas filas.",
        "La reproduccion en otro computador confirmo que el README, Compose y los secretos locales son suficientes para levantar la solucion.",
    ], size=10.0, after=4)
    add_heading(doc, "Trabajo futuro", 2)
    add_text(doc, "La evaluacion puede ampliarse con datos geograficamente diversos, medicion de memoria de mayor frecuencia, varias maquinas fisicas, pruebas de carga concurrente sobre la API y una visualizacion Leaflet. Estas extensiones no son necesarias para validar los objetivos actuales.")
    add_heading(doc, "Referencias y evidencias", 2)
    refs = [
        "Repositorio del proyecto. https://github.com/sBoterino/big-data-geoespacial",
        "Dataset en Kaggle. https://www.kaggle.com/datasets/muzammilrizvi1/motor-vehicle-collisions-crashes",
        "Evidencia de carga completa. docs/evidencias/gate4_carga_completa_2026-10-04.md",
        "Consultas con datos reales. docs/evidencias/fase5_datos_reales_2026-10-04.md",
        "Validacion Spark. docs/evidencias/fase6_logica_local_2026-10-04.md",
        "Benchmark. docs/evidencias/fase8_benchmark_2026-10-05.md y benchmark/resultados.csv",
        "Reproducibilidad. docs/evidencias/fase9_reproducibilidad_juan_2026-10-07.md",
        "MongoDB Geospatial Queries. https://www.mongodb.com/docs/manual/geospatial-queries/",
        "Docker Compose. https://docs.docker.com/compose/",
        "Apache Spark. https://spark.apache.org/docs/latest/",
    ]
    for idx, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{idx}. {ref}")
        set_run_font(r, size=8.6)

    core = doc.core_properties
    core.title = "Procesamiento y consulta de datos geoespaciales"
    core.subject = "Informe tecnico del sistema distribuido con Dask, MongoDB, Spark, Flask y Jenkins"
    core.author = "Juan Guillermo Echeverri; Sebastian Botero Velasquez; Santiago Villamizar Mejia"
    core.keywords = "Dask, Spark, MongoDB, GeoJSON, Flask, Jenkins, Docker"
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
